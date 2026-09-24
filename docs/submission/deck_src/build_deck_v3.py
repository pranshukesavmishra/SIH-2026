"""National-round deck: official SIH template + the team's approved graphics
(cropped at 300 dpi from their 10-Sept deck) + every content fix from
docs/PROJECT_STATE.md section 4. Single source of truth — rebuilds from
template.pptx each run."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from urllib.parse import quote_plus

NAVY  = RGBColor(0x0B, 0x25, 0x45)
CYAN  = RGBColor(0x0E, 0x7C, 0x99)
GREEN = RGBColor(0x1F, 0x8A, 0x5B)
GRAY  = RGBColor(0x5A, 0x64, 0x72)
LIGHT = RGBColor(0xEF, 0xF3, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED   = RGBColor(0x9C, 0x2E, 0x23)

prs = Presentation("template.pptx")
N_TESTS = 250                 # pytest tests/: 250 passed, 1 skipped
S = list(prs.slides)


def set_runs(par, runs, size=14, align=None):
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    if align is not None:
        par.alignment = align
    for text, opt in runs:
        r = par.add_run(); r.text = text
        f = r.font
        f.name = "Arial"
        f.size = Pt(opt.get("size", size))
        f.bold = opt.get("bold", False)
        f.italic = opt.get("italic", False)
        f.color.rgb = opt.get("color", NAVY)


def tb(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    return box, tf


def card(slide, x, y, w, h, fill=LIGHT, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.06
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.08)
    return shp, tf


def bullets(tf, items, size=13, gap=4, color=NAVY):
    first = True
    for item in items:
        p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        if isinstance(item, tuple):
            lead, rest = item
            set_runs(p, [(lead, {"bold": True, "size": size, "color": color}),
                         (rest, {"size": size, "color": NAVY})])
        else:
            set_runs(p, [(item, {"size": size, "color": NAVY})])


def find(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def style_pointer_box(slide, y=1.28, h=0.40, size=9):
    box = find(slide, "TextBox 8")
    box.left, box.top = Inches(1.92), Inches(max(y, 1.28))
    box.width, box.height = Inches(11.0), Inches(h)
    tf = box.text_frame
    tf.word_wrap = True
    for p in tf.paragraphs:
        text = "".join(r.text for r in p.runs)
        if not text.strip():
            p._p.getparent().remove(p._p)
            continue
        set_runs(p, [(text.strip(), {"italic": True, "size": size, "color": GRAY})])
        p.space_after = Pt(0)


def team_oval(slide):
    oval = None
    for sh in slide.shapes:
        if sh.name.startswith("Oval"):
            oval = sh
    if oval is not None:
        oval.fill.solid(); oval.fill.fore_color.rgb = NAVY
        oval.line.color.rgb = CYAN; oval.line.width = Pt(1.25)
        tf = oval.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for extra in list(tf.paragraphs[1:]):
            extra._p.getparent().remove(extra._p)
        set_runs(tf.paragraphs[0], [("ZeroDrift", {"size": 12, "bold": True, "color": WHITE})],
                 align=PP_ALIGN.CENTER)


def clean_title(slide):
    t = find(slide, "Title 1")
    if t is not None:
        for par in t.text_frame.paragraphs:
            for r in par.runs:
                r.text = r.text.replace("\x0b", "").replace("", "")


def pic(slide, path, x, y, w=None, h=None, border=None):
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)
    if border is not None:
        p.line.color.rgb = border; p.line.width = Pt(0.75)
    return p



def clone_slide(prs, source_index, insert_before):
    """
    Duplicate a slide's chrome (background bar, title placeholder, team
    oval, ISRO mark) and drop its body content, then move the copy to
    ``insert_before``.

    python-pptx has no slide-copy API, so this deep-copies the source
    slide's XML into a new slide on the same layout. Cloning rather than
    building from the blank layout is what keeps the new slide visually
    identical to the rest of the deck -- the SIH template's chrome lives
    on the slide, not on the layout, so a blank-layout slide comes out
    bare and obviously bolted on.
    """
    import copy as _copy
    src = prs.slides[source_index]
    dst = prs.slides.add_slide(src.slide_layout)
    for shp in list(dst.shapes):
        shp._element.getparent().remove(shp._element)
    keep = ("Rectangle", "Title", "Slide Number", "Footer", "Oval", "Picture")
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    for shp in src.shapes:
        if shp.name.startswith(keep):
            el = _copy.deepcopy(shp._element)
            # A copied picture still points at the SOURCE slide's image
            # relationship id; re-link it on the new slide, or the file
            # opens as corrupt in PowerPoint (LibreOffice hides this).
            for blip in el.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
                key = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                rid = blip.get(key)
                if rid:
                    blip.set(key, dst.part.relate_to(src.part.related_part(rid), RT.IMAGE))
            dst.shapes._spTree.append(el)
    ids = prs.slides._sldIdLst
    ids.insert(insert_before, ids[-1])
    return dst


def set_title(slide, text):
    t = find(slide, "Title 1")
    if t is None:
        return
    tf = t.text_frame
    tf.text = text
    set_runs(tf.paragraphs[0], [(text, {"bold": True, "size": 26, "color": NAVY})])


# ================= SLIDE 1 — TITLE =================
s = S[0]
title = find(s, "Title 7")
for par in title.text_frame.paragraphs:
    for r in par.runs:
        r.font.size = Pt(34)
title.width = Inches(10.2)
FIELDS = {
    "Problem Statement ID": "  26169",
    "Problem Statement Title": ("  AI-Based Virtual Camera Tracking System for Coarse "
                                "Alignment of Mobile FSOC Terminals"),
    "Theme": "  Smart Automation",
    "PS Category": "  Software",
    "Team ID": "  —",
    "Team Name": "  ZeroDrift",
}
box = find(s, "TextBox 9")
for p in box.text_frame.paragraphs:
    text = "".join(r.text for r in p.runs).strip()
    matched = None
    for key, val in FIELDS.items():
        if text.startswith(key):
            matched = (key, val); break
    if matched is None:
        continue
    key, val = matched
    set_runs(p, [(key + " –", {"bold": True, "size": 14, "color": NAVY}),
                 (val, {"size": 14, "color": CYAN,
                        "bold": key in ("Team Name", "Problem Statement ID")})])
    p.space_after = Pt(6)
    p.line_spacing = 1.0

shp, tf = card(s, 0.36, 5.62, 5.85, 1.10, fill=LIGHT)
set_runs(tf.paragraphs[0], [("The problem in one line:  ", {"bold": True, "size": 11.5, "color": NAVY}),
                            ("a mobile laser (FSOC) terminal must find, verify and continuously "
                             "point at its partner's beacon with hair-thin accuracy — ISRO asks "
                             "for this coarse-alignment brain, proven entirely in software.",
                             {"size": 11.5, "color": NAVY})])
pic(s, "c1_isro.jpg", 6.55, 6.32, h=0.95)
pic(s, "logo_banner.png", 8.35, 6.28, h=1.02)

# ================= SLIDE 2 — IDEA: PROBLEM → SOLUTION =================
# Left: the problem, three verified facts and ISRO's own acceptance bar.
# Right: the answer, the three-card graphic the team approved, measured
# chips, and what makes it different. Every external number has its source
# on the slide.
s = S[1]
clean_title(s); team_oval(s)
style_pointer_box(s, h=0.60, size=9)
RED_TINT = RGBColor(0xFB, 0xEE, 0xEC)
LINE = RGBColor(0xD5, 0xDD, 0xE5)

_, tf = tb(s, 0.32, 1.98, 4.95, 0.32)
set_runs(tf.paragraphs[0], [("THE PROBLEM", {"bold": True, "size": 14, "color": RED}),
                            ("   why pointing is the hard part", {"size": 10, "color": GRAY})])
PROBS = [
    ("µrad", "Laser beams are hair-thin. Both terminals must stay pointed at each "
             "other to micro-radians, on platforms that move and shake.",
     "Kaymak et al., IEEE Commun. Surveys & Tutorials 20(2), 2018"),
    ("100 Gbps", "The record 100 Gbps link to a moving drone stayed up only because a "
                 "camera tracked the target (10 Hz vision + 200 Hz tip/tilt).",
     "Walsh et al., Scientific Reports 12, 2022"),
    ("1/5 cost", "Free-space optics costs about one-fifth of laying fibre and deploys in "
                 "hours — India's last-mile and disaster links need it mobile.",
     "TEC / DoT, White Paper on Free Space Optics"),
]
y = 2.38
for big, text, src in PROBS:
    shp, _ = card(s, 0.32, y, 4.95, 0.92, fill=LIGHT)
    _, tf = tb(s, 0.44, y + 0.06, 1.30, 0.80, anchor=MSO_ANCHOR.MIDDLE)
    set_runs(tf.paragraphs[0], [(big, {"bold": True, "size": 19 if len(big) < 6 else 16, "color": RED})])
    _, tf = tb(s, 1.80, y + 0.08, 3.38, 0.60)
    set_runs(tf.paragraphs[0], [(text, {"size": 9.5, "color": NAVY})])
    _, tf = tb(s, 1.80, y + 0.68, 3.38, 0.18)
    set_runs(tf.paragraphs[0], [(src, {"italic": True, "size": 7.5, "color": GRAY})])
    y += 1.02

shp, tf = card(s, 0.32, 5.46, 4.95, 1.36, fill=RED_TINT)
set_runs(tf.paragraphs[0], [("ISRO PS26169 asks for the coarse-alignment brain:",
                             {"bold": True, "size": 10.5, "color": RED})], align=PP_ALIGN.LEFT)
p = tf.add_paragraph(); p.space_before = Pt(3)
set_runs(p, [("find the partner's beacon among stars, glints and decoys, prove it is the "
              "right one, and hold it centred — in software.", {"size": 9.5, "color": NAVY})], align=PP_ALIGN.LEFT)
for i, (v, l) in enumerate([("≤ 2 s", "acquire"), ("≤ 10 px", "error"), ("≥ 20", "FPS")]):
    x = 0.46 + i * 1.60
    shp2, tf2 = card(s, x, 6.22, 1.48, 0.50, fill=WHITE, line=RGBColor(0xE3, 0xB9, 0xB3))
    tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_runs(tf2.paragraphs[0], [(v + "  ", {"bold": True, "size": 12, "color": RED}),
                                 (l, {"size": 8.5, "color": GRAY})], align=PP_ALIGN.CENTER)

arrow = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(5.40), Inches(3.70), Inches(0.36), Inches(0.70))
arrow.fill.solid(); arrow.fill.fore_color.rgb = CYAN; arrow.line.fill.background(); arrow.shadow.inherit = False

_, tf = tb(s, 5.92, 1.98, 7.10, 0.32)
set_runs(tf.paragraphs[0], [("OUR SOLUTION — ZERODRIFT", {"bold": True, "size": 14, "color": CYAN}),
                            ("   the eyes and neck of a laser terminal", {"size": 10, "color": GRAY})])
pic(s, "c2_cards.jpg", 5.92, 2.38, w=4.70)                        # h = 2.36
for i, chp in enumerate(["c2_chip1p.jpg", "c2_chip2p.jpg", "c2_chip3.jpg"]):
    pic(s, chp, 10.78, 2.38 + i * 0.81, h=0.74)
_, tf = tb(s, 10.78, 4.80, 2.25, 0.18)
set_runs(tf.paragraphs[0], [("64 simulated LEO passes", {"italic": True, "size": 7.5, "color": GRAY})],
         align=PP_ALIGN.CENTER)

DIFF = [
    ("Identity, not brightness", "Locks only on the 4 Hz blink. Lamps, stars, glints and "
                                 "wrong-rate strobes are rejected."),
    ("Built and running", "Live webcam demo, MK1 servo rig, MK2 stepper terminal "
                          "— ₹6,709 bill of materials."),
    ("Measured, not claimed", f"{N_TESTS} passing tests. Every number on these slides "
                              "regenerates from one command."),
]
for i, (h, t) in enumerate(DIFF):
    x = 5.92 + i * 2.39
    shp, tf = card(s, x, 5.08, 2.28, 1.74, fill=LIGHT, line=LINE)
    tf.margin_top = Inches(0.12)
    set_runs(tf.paragraphs[0], [(f"0{i + 1}", {"bold": True, "size": 18, "color": CYAN})], align=PP_ALIGN.LEFT)
    p = tf.add_paragraph(); p.space_before = Pt(2)
    set_runs(p, [(h, {"bold": True, "size": 11, "color": NAVY})], align=PP_ALIGN.LEFT)
    p = tf.add_paragraph(); p.space_before = Pt(3)
    set_runs(p, [(t, {"size": 9.5, "color": NAVY})], align=PP_ALIGN.LEFT)

# ================= SLIDE 3 — TECHNICAL: THE LOOP, THE STATES, THE NUMBERS =================
s = S[2]
clean_title(s); team_oval(s)
style_pointer_box(s, h=0.26, size=9)
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.oxml.ns import qn
import json as _json
from lxml import etree

STAGES = [
    ("Camera", "640×480 · 30 fps", "rides the gimbal, boresighted"),
    ("Detect", "salt & pepper filter", "sub-pixel spots · CFAR"),
    ("Identify", "4 Hz blink test", "decoys + wrong rates vetoed"),
    ("Track", "IMM Kalman filter", "coasts the dark half-cycle"),
    ("Control", "Smith predictor", "type-2 loop cancels drift"),
    ("Gimbal", "2× NEMA17 steppers", "AS5600 angle feedback"),
]
BW, BH, GAP, X0, Y0 = 1.86, 1.02, 0.26, 0.42, 1.76
for i, (h, l1, l2) in enumerate(STAGES):
    x = X0 + i * (BW + GAP)
    shp, tf = card(s, x, Y0, BW, BH, fill=LIGHT, line=RGBColor(0xB9, 0xD3, 0xDE))
    tf.margin_left = Inches(0.10); tf.margin_top = Inches(0.04); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_runs(tf.paragraphs[0], [(f"{i + 1}  ", {"bold": True, "size": 12, "color": CYAN}),
                                (h, {"bold": True, "size": 14, "color": NAVY})])
    p = tf.add_paragraph(); p.space_before = Pt(4)
    set_runs(p, [(l1, {"bold": True, "size": 10, "color": NAVY})])
    p = tf.add_paragraph(); p.space_before = Pt(1)
    set_runs(p, [(l2, {"size": 9, "color": GRAY})])
    if i < len(STAGES) - 1:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + BW + 0.03), Inches(Y0 + BH / 2 - 0.10),
                                Inches(GAP - 0.06), Inches(0.20))
        ar.fill.solid(); ar.fill.fore_color.rgb = CYAN; ar.line.fill.background(); ar.shadow.inherit = False


def seg(slide, x1, y1, x2, y2, head=False):
    c = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = CYAN; c.line.width = Pt(1.75)
    if head:
        ln = c.line._get_or_add_ln()
        te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    return c


# the closed loop: gimbal back to camera, under the chain
yl = Y0 + BH + 0.20
xg = X0 + 5 * (BW + GAP) + BW / 2; xc = X0 + BW / 2
seg(s, xg, Y0 + BH, xg, yl); seg(s, xg, yl, xc, yl); seg(s, xc, yl, xc, Y0 + BH, head=True)
_, tf = tb(s, 3.2, yl + 0.03, 7.0, 0.22)
set_runs(tf.paragraphs[0], [("closed optical loop — the camera moves with the laser, so the error is "
                             "measured on the image itself", {"italic": True, "size": 9, "color": CYAN})],
         align=PP_ALIGN.CENTER)

# acquisition state machine
ys = 3.50
_, tf = tb(s, 0.42, ys + 0.06, 2.1, 0.26)
set_runs(tf.paragraphs[0], [("STATE MACHINE", {"bold": True, "size": 10, "color": GRAY})])
PILLS = [("SEARCH", RGBColor(0xE3, 0xEB, 0xFB)), ("ACQUIRE", RGBColor(0xFD, 0xF1, 0xD8)),
         ("TRACK", RGBColor(0xDF, 0xF3, 0xE8)), ("COAST", RGBColor(0xFD, 0xE8, 0xD8)),
         ("REACQUIRE", RGBColor(0xFB, 0xE3, 0xE1))]
PW = 1.55
for i, (lab, col) in enumerate(PILLS):
    x = 2.30 + i * (PW + 0.44)
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(ys), Inches(PW), Inches(0.38))
    shp.adjustments[0] = 0.5; shp.fill.solid(); shp.fill.fore_color.rgb = col
    shp.line.fill.background(); shp.shadow.inherit = False
    tf = shp.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_runs(tf.paragraphs[0], [(lab, {"bold": True, "size": 10, "color": NAVY})], align=PP_ALIGN.CENTER)
    if i < len(PILLS) - 1:
        _, tf = tb(s, x + PW, ys + 0.02, 0.44, 0.34, anchor=MSO_ANCHOR.MIDDLE)
        set_runs(tf.paragraphs[0], [("⇄" if lab == "TRACK" else "→", {"bold": True, "size": 14, "color": CYAN})],
                 align=PP_ALIGN.CENTER)

# ---- chart A: pointing error on the ISS pass, native, log scale ----
run = _json.load(open("../../media/telemetry_run.json"))
fr = run["frames"]
step = max(1, len(fr) // 180)
pts = [(f["t"], f["err_urad"]) for f in fr[::step] if f.get("err_urad")]
fov = run["constants"]["fov_half_width_urad"]
cd = XyChartData()
s1 = cd.add_series("Pointing error")
for t_, e_ in pts: s1.add_data_point(t_, max(e_, 10))
s2 = cd.add_series("Field-of-view edge")
s2.add_data_point(0, fov); s2.add_data_point(pts[-1][0], fov)
gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(0.32), Inches(4.10),
                        Inches(4.95), Inches(2.72), cd)
ch = gf.chart
ch.has_title = True
ch.chart_title.text_frame.text = "Pointing error, ISS pass (µrad, log)"
tp = ch.chart_title.text_frame.paragraphs[0]
set_runs(tp, [("ISS pass: ", {"bold": True, "size": 10.5, "color": NAVY}),
              (f"lock in {run['summary']['acquisition_time_s']:.2f} s, median "
               f"{run['summary']['pointing_error_urad']['p50']:.0f} µrad", {"size": 10, "color": NAVY})])
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
ch.legend.font.size = Pt(8); ch.legend.font.color.rgb = GRAY
ser = ch.plots[0].series
ser[0].format.line.color.rgb = CYAN; ser[0].format.line.width = Pt(1.25); ser[0].smooth = False
ser[1].format.line.color.rgb = RED; ser[1].format.line.width = Pt(1.25); ser[1].format.line.dash_style = 4
va, ca = ch.value_axis, ch.category_axis
va.minimum_scale, va.maximum_scale = 10, 100000
sc = va._element.find(qn("c:scaling"))
lb = etree.SubElement(sc, qn("c:logBase")); lb.set("val", "10"); sc.remove(lb); sc.insert(0, lb)
va.has_major_gridlines = True; va.major_gridlines.format.line.color.rgb = RGBColor(0xE4, 0xEA, 0xF0)
va.tick_labels.font.size = Pt(8); va.tick_labels.font.color.rgb = GRAY; va.tick_labels.number_format = '#,##0'
va.tick_labels.number_format_is_linked = False
ca.minimum_scale, ca.maximum_scale = 0, 45
ca.tick_labels.font.size = Pt(8); ca.tick_labels.font.color.rgb = GRAY
ca.has_title = True; ca.axis_title.text_frame.text = "time (s)"
set_runs(ca.axis_title.text_frame.paragraphs[0], [("time (s)", {"size": 8, "color": GRAY})])
for ax in (va, ca): ax.format.line.color.rgb = RGBColor(0xC8, 0xD2, 0xDC)

# ---- chart B: the live-demo tracker, measured on 12 synthetic scenes ----
BENCH = [("Still", 100.0), ("Slow circle", 100.0), ("Figure-8", 100.0), ("Fast figure-8", 93.7),
         ("Hand, run 1", 95.4), ("Hand, run 2", 99.5), ("Hand, run 3", 97.9), ("Past a lamp", 100.0),
         ("Phone screen", 100.0), ("4.6 Hz beacon", 99.6), ("Torch turning", 100.0), ("Wrong-rate strobes", 99.7)]
cd = CategoryChartData()
cd.categories = [b[0] for b in BENCH][::-1]
cd.add_series("Lock held (%)", [b[1] for b in BENCH][::-1])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(5.45), Inches(4.10), Inches(4.35), Inches(2.72), cd)
ch = gf.chart
ch.has_title = True; ch.chart_title.text_frame.text = "Live demo: lock held"
set_runs(ch.chart_title.text_frame.paragraphs[0],
         [("Live demo: ", {"bold": True, "size": 10.5, "color": NAVY}),
          ("% of time locked, 12 scenes", {"size": 10, "color": NAVY})])
ch.has_legend = False
pl = ch.plots[0]; pl.gap_width = 45
pl.series[0].format.fill.solid(); pl.series[0].format.fill.fore_color.rgb = GREEN
pl.has_data_labels = True
dl = pl.data_labels; dl.number_format = '0.0'; dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.INSIDE_END; dl.font.size = Pt(7.5); dl.font.color.rgb = WHITE; dl.font.bold = True
va, ca = ch.value_axis, ch.category_axis
va.minimum_scale, va.maximum_scale = 0, 100
va.has_major_gridlines = False; va.visible = False
ca.tick_labels.font.size = Pt(8); ca.tick_labels.font.color.rgb = NAVY
ca.format.line.color.rgb = RGBColor(0xC8, 0xD2, 0xDC)

# ---- the stack ----
shp, tf = card(s, 9.98, 4.10, 3.03, 2.72, fill=LIGHT, line=LINE)
tf.margin_left = Inches(0.14); tf.margin_top = Inches(0.10)
set_runs(tf.paragraphs[0], [("BUILT WITH", {"bold": True, "size": 10, "color": GRAY})], align=PP_ALIGN.LEFT)
STACK = [("Engine", "Python · NumPy · OpenCV"), ("Web demo", "JavaScript · three.js · Web Serial"),
         ("Hardware", "Arduino Nano · A4988 · AS5600"), ("Proof", f"pytest · {N_TESTS} tests · GitHub Actions"),
         ("Acquisition", "0.55–1.1 s in the live demo")]
for h, t in STACK:
    p = tf.add_paragraph(); p.space_before = Pt(5)
    set_runs(p, [(h, {"bold": True, "size": 10, "color": NAVY})], align=PP_ALIGN.LEFT)
    p = tf.add_paragraph()
    set_runs(p, [(t, {"size": 9.5, "color": GRAY})], align=PP_ALIGN.LEFT)

# ================= SLIDE 4 — FEASIBILITY =================
s = S[3]
clean_title(s); team_oval(s)
style_pointer_box(s, h=0.26, size=9)
pic(s, "c4_cols.jpg", 1.05, 1.62, w=11.22)                        # h = 4.32
pic(s, "c4_strip.jpg", 3.42, 6.00, w=6.50)                        # h = 0.93
_, tf = tb(s, 0.35, 6.06, 2.95, 0.80, anchor=MSO_ANCHOR.MIDDLE)
bullets(tf, [("64 runs, our LEO scenarios →  ",
              "never a wrong lock, never a failed acquisition.")], size=10.5, gap=0)
_, tf = tb(s, 10.05, 6.06, 2.95, 0.80, anchor=MSO_ANCHOR.MIDDLE)
bullets(tf, [("Reproducible:  ",
              "every number regenerates from one command.")], size=10.5, gap=0)

# ================= SLIDE 5 — IMPACT =================
s = S[4]
clean_title(s); team_oval(s)
style_pointer_box(s, h=0.26, size=9)

_, tf = tb(s, 0.30, 1.64, 4.25, 0.42)
set_runs(tf.paragraphs[0], [("ECONOMIC FEASIBILITY", {"bold": True, "size": 13.5, "color": GREEN}),
                            ("  — the arithmetic that removes the cost barrier",
                             {"size": 10.5, "color": GRAY})])
pic(s, "c_econ_capital.jpg", 0.30, 2.06, w=4.06, border=RGBColor(0xD5, 0xDD, 0xE5))  # h 2.11
pic(s, "c_econ_stats.jpg", 0.30, 4.26, w=4.06, border=RGBColor(0xD5, 0xDD, 0xE5))    # h 1.30

pic(s, "c5_cards.jpg", 4.66, 1.62, w=8.20)                        # h = 3.96
pic(s, "c5_audience.jpg", 0.35, 5.66, w=5.72)                     # h = 1.25
pic(s, "c5_linkbox.png", 7.20, 5.66, w=5.75)                      # h = 1.26


# ============ NEW SLIDE — PS26169 COMPLIANCE ============
# Placed before References, while the evaluators are still in technical
# mode. Technical Evaluation is 20% and its first listed criterion is
# "Understanding of the problem"; a table mapping their own rows to our
# values demonstrates that in one glance, where restating their
# background paragraph demonstrates nothing.
s = clone_slide(prs, 2, 5)
set_title(s, "PS26169 compliance: their table, our values")
_t = find(s, "Title 1")
if _t is not None:   # clear of the SIH logo on the right
    _t.left, _t.width = Inches(1.95), Inches(8.60)
    _t.top, _t.height = Inches(0.30), Inches(0.80)
    _t.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _t.text_frame.word_wrap = True
    for _r in _t.text_frame.paragraphs[0].runs:
        _r.font.size = Pt(24)
# No pointer box here: the clone deliberately drops the template's body
# text placeholder, and this slide is the table.

ROWS = [
    ("Camera resolution",        "640 x 480",                 "640 x 480",                    True),
    ("Camera FOV",               "4\u00b0 x 3\u00b0 default",          "4\u00b0 x 3\u00b0",                     True),
    ("Camera update rate",       "\u2265 30 Hz",                  "30 Hz",                        True),
    ("Max pan / tilt speed",     "5-10\u00b0/s, default 5",       "5\u00b0/s",                       True),
    ("Target shape / size",      "square, 10 x 10 px",        "square, 10 x 10 px",           True),
    ("Motion patterns",          "\u2265 4, incl. Figure-of-8",   "7, incl. Figure-of-8",         True),
    ("Image noise",              "S&P ~10%, Gaussian, Poisson", "all three, S&P at 10%",      True),
    ("Camera jitter",            "\u00b1 20 px / frame",          "19.8 px p99, measured",        True),
    ("Platform motion",          "\u00b1 20 px / frame, linear",  "linear drift, 6 px / frame",   True),
    ("Atmospheric conditions",   "Clear / Haze / Fog / Rain / Low light", "all five",          True),
    ("MP4 input, bypassing PTZ", "required (Benchmark-2)",    "--video, working",             True),
    ("Acquisition time",         "\u2264 2 s",                    "25.3 s \u2192 3.53 s  \u2014 open",  False),
    ("Tracking error (median)",  "\u2264 10 px",                  "2,367 \u2192 170 px  \u2014 open",  False),
    ("Processing speed",         "\u2265 20 FPS",                 "9.2 \u2192 14.7 FPS  \u2014 open",  False),
]

x0, y0, w, rowh = 0.62, 1.66, 12.1, 0.325
cw = (4.35, 3.85, 3.55, 0.35)

hdr = card(s, x0, y0, w, rowh, fill=NAVY)
cx = x0 + 0.12
for txt, width in zip(("PS26169 PARAMETER", "SPECIFIED", "ZERODRIFT", ""), cw):
    _, tf = tb(s, cx, y0, width, rowh, anchor=MSO_ANCHOR.MIDDLE)
    set_runs(tf.paragraphs[0], [(txt, {"bold": True, "size": 10, "color": WHITE})])
    cx += width

for i, (name, spec, ours, ok) in enumerate(ROWS):
    y = y0 + rowh * (i + 1)
    if ok is False:
        card(s, x0, y, w, rowh, fill=RGBColor(0xFC, 0xF0, 0xEE))
    elif ok is None:
        card(s, x0, y, w, rowh, fill=RGBColor(0xFD, 0xF8, 0xE8))
    elif i % 2 == 0:
        card(s, x0, y, w, rowh, fill=LIGHT)
    cx = x0 + 0.12
    colour = RED if ok is False else NAVY
    for txt, width, bold in ((name, cw[0], False), (spec, cw[1], False),
                             (ours, cw[2], True)):
        _, tf = tb(s, cx, y, width, rowh, anchor=MSO_ANCHOR.MIDDLE)
        set_runs(tf.paragraphs[0], [(txt, {"size": 9.5, "bold": bold, "color": colour})])
        cx += width
    mark = "\u2717" if ok is False else ("~" if ok is None else "\u2713")
    _, tf = tb(s, cx, y, cw[3], rowh, anchor=MSO_ANCHOR.MIDDLE)
    set_runs(tf.paragraphs[0], [(mark, {"size": 11, "bold": True,
                                        "color": RED if ok is False else
                                        (GRAY if ok is None else GREEN)})],
             align=PP_ALIGN.CENTER)

# The honest footnote. A compliance table with a number we cannot
# reproduce on demand is worse than no table -- Q&A is 20% of the
# technical evaluation and that is the row they will ask us to run.
foot_y = y0 + rowh * (len(ROWS) + 1) + 0.10
card(s, x0, foot_y, w, 0.76, fill=RGBColor(0xF3, 0xF7, 0xFA))
_, tf = tb(s, x0 + 0.16, foot_y + 0.02, w - 0.32, 0.72)
bullets(tf, [
    ("The three open rows are open, not footnoted.  ",
     "Arrows are before \u2192 after our own fix, both measured with the full specified "
     "disturbance present. A constant-velocity platform drift is a ramp, and a PD loop "
     "has finite steady-state error to a ramp \u2014 wrong loop order, not bad tuning. "
     "Adding a second integrator (type 2, leaky) cut acquisition 7x and median error "
     "14x. It is implemented and shipped; it is not yet enough."),
    ("Every figure regenerates from  ",
     "scenarios/ps26169_benchmark.yaml, which encodes their table and nothing else."),
], size=9, gap=2)

# ================= SLIDE 6 — REFERENCES =================
s = S[5]
clean_title(s); team_oval(s)
style_pointer_box(s, h=0.22, size=9)


def scholar(t):
    return "https://scholar.google.com/scholar?q=" + quote_plus('"' + t + '"')


_, tf = tb(s, 0.42, 1.58, 6.0, 0.34)
set_runs(tf.paragraphs[0], [("Research foundations", {"bold": True, "size": 14, "color": NAVY})])
REFS = [
    ("Kaymak et al.,  ", "“A Survey on Acquisition, Tracking and Pointing Mechanisms for Mobile FSO,” IEEE Comm. Surveys & Tutorials, 2018.",
     scholar("A Survey on Acquisition Tracking and Pointing Mechanisms for Mobile Free-Space Optical Communications")),
    ("Kaushal & Kaddoum,  ", "“Optical Communication in Space: Challenges and Mitigation Techniques,” IEEE CS&T, 2017.",
     scholar("Optical Communication in Space Challenges and Mitigation Techniques")),
    ("Blackman & Popoli,  ", "Design and Analysis of Modern Tracking Systems — IMM filtering, CFAR detection.",
     scholar("Design and Analysis of Modern Tracking Systems Blackman Popoli")),
    ("Bar-Shalom et al.,  ", "Estimation with Applications to Tracking and Navigation — Kalman/IMM theory.",
     scholar("Estimation with Applications to Tracking and Navigation Bar-Shalom")),
    ("Smith, O.J.M.,  ", "“Closer Control of Loops with Dead Time” — the Smith predictor, 1957.",
     scholar("Closer Control of Loops with Dead Time Smith")),
    ("Vallado & Crawford,  ", "“SGP4 Orbit Determination” — AIAA 2008; real ISS TLE propagation.",
     scholar("SGP4 Orbit Determination Vallado Crawford")),
    ("Andrews & Phillips,  ", "Laser Beam Propagation through Random Media — turbulence & scintillation.",
     scholar("Laser Beam Propagation through Random Media Andrews Phillips")),
]
_, tf = tb(s, 0.42, 1.98, 6.15, 3.45)
first = True
for lead, rest, url in REFS:
    pr = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    pr.space_after = Pt(4)
    set_runs(pr, [(lead, {"bold": True, "size": 10, "color": NAVY}),
                  (rest + "  ", {"size": 10, "color": NAVY})])
    lk = pr.add_run(); lk.text = "[link]"
    lk.font.name = "Arial"; lk.font.size = Pt(9.5)
    lk.font.color.rgb = CYAN; lk.font.underline = True
    lk.hyperlink.address = url
pic(s, "c6_logos.jpg", 1.10, 5.52, w=4.40)                        # h = 1.38

_, tf = tb(s, 6.95, 1.58, 6.0, 0.34)
set_runs(tf.paragraphs[0], [("Our work — open and verifiable", {"bold": True, "size": 14, "color": NAVY})])
_, tf = tb(s, 6.95, 1.98, 6.0, 1.42)
pr = tf.paragraphs[0]
set_runs(pr, [("Repository:  ", {"bold": True, "size": 11, "color": NAVY})])
lk = pr.add_run(); lk.text = "github.com/pranshukesavmishra/SIH-2026"
lk.font.name = "Arial"; lk.font.size = Pt(11)
lk.font.color.rgb = CYAN; lk.font.underline = True
lk.hyperlink.address = "https://github.com/pranshukesavmishra/SIH-2026"
pr.space_after = Pt(4)
pr2 = tf.add_paragraph()
set_runs(pr2, [("Live demo (replay console):  ", {"bold": True, "size": 11, "color": NAVY})])
lk = pr2.add_run(); lk.text = "zerodrift-fsoc-pat.netlify.app"
lk.font.name = "Arial"; lk.font.size = Pt(11)
lk.font.color.rgb = CYAN; lk.font.underline = True
lk.hyperlink.address = "https://zerodrift-fsoc-pat.netlify.app/"
pr2.space_after = Pt(4)
bullets(tf, [
    ("Inside:  ", f"5,500+ documented lines · {N_TESTS} passing tests · 6 validated scenarios · Monte-Carlo logs · technical report · user manual · demo video."),
    ("Reproducible:  ", "every figure regenerates from one command; every random draw is logged and replayable."),
], size=10.5, gap=4)

_, tf = tb(s, 6.95, 3.52, 5.5, 0.26)
set_runs(tf.paragraphs[0], [("Compared with existing approaches", {"bold": True, "size": 11.5, "color": NAVY})])
pic(s, "c6_table.jpg", 6.95, 3.80, w=5.30)                        # h = 1.86

shp, tf = card(s, 6.95, 5.74, 6.0, 1.16, fill=NAVY)
set_runs(tf.paragraphs[0], [("Why ZeroDrift wins on this problem",
                             {"bold": True, "size": 13, "color": WHITE})])
for line in ["All six mandatory deliverables already exist today",
             "Judged numbers, honestly measured against hidden ground truth",
             "AI where it earns its place — benchmarked against classical methods"]:
    p = tf.add_paragraph()
    set_runs(p, [("✓  ", {"bold": True, "size": 10.5, "color": RGBColor(0x7F, 0xD4, 0xA8)}),
                 (line, {"size": 10.5, "color": WHITE})])
    p.space_before = Pt(2)

# ============ delete the template's instructions slide ============
# Found by its content, not by index. It used to be removed as slide 6,
# which silently became the wrong slide the moment anything was inserted
# ahead of it -- and inserting the compliance slide did exactly that.
xml_slides = prs.slides._sldIdLst
_marker = "instruction"
_victims = [i for i, sl in enumerate(prs.slides)
            if any(sh.has_text_frame and _marker in sh.text_frame.text.lower()
                   for sh in sl.shapes)]
if len(_victims) != 1:
    raise SystemExit(f"expected exactly one instructions slide, found {_victims}")
xml_slides.remove(list(xml_slides)[_victims[0]])

for _sl in prs.slides:
    for _sh in _sl.shapes:
        if _sh.has_text_frame and "@SIH Idea submission" in _sh.text_frame.text:
            for _par in _sh.text_frame.paragraphs:
                for _r in _par.runs:
                    if "@SIH Idea submission" in _r.text:
                        _r.text = "Team ZeroDrift · SIH 2026 · PS 26169"

core = prs.core_properties
core.title = "ZeroDrift — SIH 2026 Idea Submission (PS 26169, ISRO)"
core.author = "Team ZeroDrift, Jabalpur Engineering College"
core.subject = "AI-Based Virtual Camera Tracking for Coarse Alignment of Mobile FSOC Terminals"
core.keywords = "SIH 2026, PS 26169, ISRO, FSOC, ZeroDrift"
prs.save("ZeroDrift_SIH26169.pptx")
print("saved ZeroDrift_SIH26169.pptx")
